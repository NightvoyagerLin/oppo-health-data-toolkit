# SPDX-License-Identifier: MIT
"""
OPPO 健康数据导出工具（GUI 版）
================================
功能：
- 选择解密后的 SQLite 数据库文件
- 显示数据库信息（表数量、数据量、时间范围）
- 选择导出时间范围（最近7天/30天/90天/全部/自定义）
- 选择导出格式（CSV/JSON/SQLite/Excel）
- 选择要导出的表
- 一键导出

使用方法：
    python health_export_gui.py
"""

import os
import sys          # L11 修复后重新需要：__main__ 用 sys.exit(main()) 传递退出码

# 强制 UTF-8 输出，避免 Windows 控制台/重定向乱码（V1.1 修复）
try:
    from utils import ensure_utf8_stdout
    ensure_utf8_stdout()
except Exception:
    pass

import json
import csv
import sqlite3
# `root = tk.Tk()`。在无图形库或无可显示环境（服务器、CI、容器）下，前者抛
# ImportError、后者抛 TclError，连"本文件能否被导入"都做不到。
# 现改为容错导入 + main() 内友好提示，使本模块在无显示环境下仍可被安全导入
# （便于静态检查与自动化测试），真正需要开界面时再给出可执行的替代方案。
try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    _TK_AVAILABLE = True
    _TK_ERROR = None
except Exception as _e:                # ImportError（无 tkinter）等
    tk = ttk = filedialog = messagebox = None
    _TK_AVAILABLE = False
    _TK_ERROR = _e
# 引入 timezone，日期展示统一按中国时区(UTC+8)，
# 与 split_db_by_day.py 的 CN 口径一致（此前用本地时区，跨时区机器会偏移）。
from datetime import datetime, timedelta, timezone
import threading


def q_ident(name):
    """把表名/列名安全地包成 SQLite 标识符。
    多处用 f'"{table}"' 直接拼 SQL，名字里含双引号会提前闭合标识符、破坏语句
    （SQL 注入风险）——按 SQL 规范把 `"` 翻倍即可安全嵌入。
    与 json_to_sqlite.q_ident、mcp_server._check_ident 同一标准。"""
    return '"' + str(name).replace('"', '""').replace("\x00", "") + '"'


# 事件时间列清单：与 split_db_by_day 共用 **utils.EVENT_TIME_COLUMNS** 这一唯一来源。
# 导入失败即抛出明确错误，不静默降级为旧口径。
from utils import EVENT_TIME_COLUMNS, atomic_write, safe_fs_name
CN = timezone(timedelta(hours=8))


def _norm_ymd(v):
    """把 date 列的 YYYYMMDD 整数/字符串规整为 YYYY-MM-DD；已是 YYYY-MM-DD 则原样返回。
    统一归一后再比较，避免与时间戳列格式混排导致界面时间范围显示错乱。
    """
    s = str(v).strip()
    if len(s) == 8 and s.isdigit():
        return "%s-%s-%s" % (s[0:4], s[4:6], s[6:8])
    if len(s) == 10 and s[4] == "-" and s[7] == "-":
        return s
    return s


class HealthExportApp:
    def __init__(self, root):
        self.root = root
        self.root.title("OPPO 健康数据导出工具")
        self.root.geometry("900x700")
        self.root.minsize(800, 600)

        # 变量
        self.db_path = tk.StringVar()
        self.db_info = {}
        self.tables = []
        self.table_vars = {}
        self.time_range = tk.StringVar(value="30")  # 默认最近30天
        self.export_format = tk.StringVar(value="csv")
        self.custom_start = tk.StringVar()
        self.custom_end = tk.StringVar()
        self.export_dir = tk.StringVar()

        # 创建界面
        self.create_widgets()

    def create_widgets(self):
        # 主容器
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)

        # ========== 1. 数据库选择 ==========
        db_frame = ttk.LabelFrame(main_frame, text="1. 选择数据库文件", padding="10")
        db_frame.pack(fill=tk.X, pady=(0, 10))

        db_row = ttk.Frame(db_frame)
        db_row.pack(fill=tk.X)

        ttk.Label(db_row, text="数据库文件:").pack(side=tk.LEFT)
        ttk.Entry(db_row, textvariable=self.db_path, width=60).pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Button(db_row, text="浏览...", command=self.browse_db).pack(side=tk.LEFT)
        ttk.Button(db_row, text="加载", command=self.load_db).pack(side=tk.LEFT, padx=5)

        # 数据库信息显示
        self.info_text = tk.Text(db_frame, height=4, wrap=tk.WORD, state=tk.DISABLED)
        self.info_text.pack(fill=tk.X, pady=(10, 0))

        # ========== 2. 时间范围选择 ==========
        time_frame = ttk.LabelFrame(main_frame, text="2. 选择导出时间范围", padding="10")
        time_frame.pack(fill=tk.X, pady=(0, 10))

        # 预设选项
        preset_row = ttk.Frame(time_frame)
        preset_row.pack(fill=tk.X, pady=(0, 5))

        ttk.Label(preset_row, text="快速选择:").pack(side=tk.LEFT)

        time_options = [
            ("最近7天", "7"),
            ("最近30天", "30"),
            ("最近90天", "90"),
            ("最近180天", "180"),
            ("最近1年", "365"),
            ("全部数据", "all"),
        ]

        for text, value in time_options:
            ttk.Radiobutton(preset_row, text=text, variable=self.time_range,
                           value=value, command=self.on_time_change).pack(side=tk.LEFT, padx=5)

        # 自定义日期
        custom_row = ttk.Frame(time_frame)
        custom_row.pack(fill=tk.X, pady=(5, 0))

        ttk.Radiobutton(custom_row, text="自定义:", variable=self.time_range,
                       value="custom", command=self.on_time_change).pack(side=tk.LEFT)

        ttk.Label(custom_row, text="开始日期:").pack(side=tk.LEFT, padx=(10, 5))
        ttk.Entry(custom_row, textvariable=self.custom_start, width=12).pack(side=tk.LEFT)
        ttk.Label(custom_row, text="(格式: 20260901)").pack(side=tk.LEFT, padx=5)

        ttk.Label(custom_row, text="结束日期:").pack(side=tk.LEFT, padx=(10, 5))
        ttk.Entry(custom_row, textvariable=self.custom_end, width=12).pack(side=tk.LEFT)
        ttk.Label(custom_row, text="(格式: 20260907)").pack(side=tk.LEFT, padx=5)

        # ========== 3. 导出格式选择 ==========
        format_frame = ttk.LabelFrame(main_frame, text="3. 选择导出格式", padding="10")
        format_frame.pack(fill=tk.X, pady=(0, 10))

        format_row = ttk.Frame(format_frame)
        format_row.pack(fill=tk.X)

        formats = [
            ("CSV（每表一个CSV文件）", "csv"),
            ("JSON（一个JSON文件）", "json"),
            ("SQLite（一个.db文件）", "sqlite"),
            ("Excel（一个.xlsx文件，每表一个sheet）", "excel"),
        ]

        for text, value in formats:
            ttk.Radiobutton(format_row, text=text, variable=self.export_format,
                           value=value).pack(side=tk.LEFT, padx=10)

        # ========== 4. 表选择 ==========
        table_frame = ttk.LabelFrame(main_frame, text="4. 选择要导出的表", padding="10")
        table_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        # 按钮行
        btn_row = ttk.Frame(table_frame)
        btn_row.pack(fill=tk.X, pady=(0, 5))

        ttk.Button(btn_row, text="全选", command=self.select_all_tables).pack(side=tk.LEFT)
        ttk.Button(btn_row, text="全不选", command=self.deselect_all_tables).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_row, text="反选", command=self.invert_tables).pack(side=tk.LEFT)
        ttk.Label(btn_row, text="（已选择 0 张表）").pack(side=tk.RIGHT)
        self.selected_count_label = btn_row.winfo_children()[-1]

        # 表列表（带滚动条）
        list_frame = ttk.Frame(table_frame)
        list_frame.pack(fill=tk.BOTH, expand=True)

        self.table_canvas = tk.Canvas(list_frame)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.table_canvas.yview)
        self.table_canvas.configure(yscrollcommand=scrollbar.set)

        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.table_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.table_inner = ttk.Frame(self.table_canvas)
        self.table_canvas.create_window((0, 0), window=self.table_inner, anchor="nw")
        self.table_inner.bind("<Configure>", lambda e: self.table_canvas.configure(scrollregion=self.table_canvas.bbox("all")))

        # ========== 5. 导出按钮 ==========
        export_frame = ttk.Frame(main_frame)
        export_frame.pack(fill=tk.X, pady=(0, 5))

        ttk.Label(export_frame, text="导出到:").pack(side=tk.LEFT)
        ttk.Entry(export_frame, textvariable=self.export_dir, width=50).pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
        ttk.Button(export_frame, text="浏览...", command=self.browse_export_dir).pack(side=tk.LEFT)

        self.export_btn = ttk.Button(export_frame, text="开始导出", command=self.start_export)
        self.export_btn.pack(side=tk.LEFT, padx=10)

        # 进度条
        self.progress = ttk.Progressbar(main_frame, mode="determinate")
        self.progress.pack(fill=tk.X, pady=(0, 5))

        # 日志输出
        log_frame = ttk.LabelFrame(main_frame, text="日志", padding="5")
        log_frame.pack(fill=tk.BOTH, expand=True)

        self.log_text = tk.Text(log_frame, height=8, wrap=tk.WORD, state=tk.DISABLED)
        log_scrollbar = ttk.Scrollbar(log_frame, orient="vertical", command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=log_scrollbar.set)

        log_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    def _safe_after(self, ms, fn):
        """G1补充：root 销毁（用户在导出中关窗）后 after 会抛 TclError；
        统一经此包装吞掉，避免在 except/finally 中抛出时掩盖业务异常。"""
        try:
            self.root.after(ms, fn)
        except Exception:
            pass

    def log(self, msg):
        """输出日志。do_export 运行在工作线程，而 Tkinter 控件不允许跨线程直接操作
        （轻则界面异常，重则解释器崩溃）。此处按调用线程分流：主线程直接写，
        工作线程经 root.after 调度回主线程执行。"""
        if threading.current_thread() is threading.main_thread():
            self._log_impl(msg)
        else:
            self._safe_after(0, lambda: self._log_impl(msg))

    def _log_impl(self, msg):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, f"[{timestamp}] {msg}\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)

    def browse_db(self):
        """浏览数据库文件"""
        path = filedialog.askopenfilename(
            title="选择解密后的 SQLite 数据库文件",
            filetypes=[("SQLite 数据库", "*.db"), ("所有文件", "*.*")]
        )
        if path:
            self.db_path.set(path)

    def browse_export_dir(self):
        """浏览导出目录"""
        path = filedialog.askdirectory(title="选择导出目录")
        if path:
            self.export_dir.set(path)

    def load_db(self):
        """加载数据库"""
        path = self.db_path.get()
        if not path or not os.path.exists(path):
            messagebox.showerror("错误", "请选择有效的数据库文件！")
            return

        conn = None  # G2修复：预置 None，保证 finally 中可安全关闭
        try:
            conn = sqlite3.connect(path)
            cursor = conn.cursor()

            # 获取所有表
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")
            self.tables = [row[0] for row in cursor.fetchall()]

            # 过滤掉内部表
            self.tables = [t for t in self.tables if t not in ("sqlite_sequence", "room_master_table")]

            # 统计信息
            total_rows = 0
            table_info = []
            min_date = None
            max_date = None

            for table in self.tables:
                try:
                    cursor.execute(f'SELECT COUNT(*) FROM {q_ident(table)}')
                    count = cursor.fetchone()[0]
                    total_rows += count
                    table_info.append((table, count))

                    # 尝试获取日期范围
                    # 候选列同样改用共享清单（与 get_date_columns 一致）
                    for date_col in ["date"] + list(EVENT_TIME_COLUMNS):
                        try:
                            cursor.execute(f'SELECT MIN({q_ident(date_col)}), MAX({q_ident(date_col)}) FROM {q_ident(table)} WHERE {q_ident(date_col)} IS NOT NULL')
                            row = cursor.fetchone()
                            if row and row[0] is not None:
                                if date_col == "date":
                                    # MIN/MAX 可能为 NULL（缺值行），分开判空
                                    # date 列规整为 YYYY-MM-DD，与下方时间戳列格式统一
                                    d_min = _norm_ymd(row[0]) if row[0] is not None else None
                                    d_max = _norm_ymd(row[1]) if row[1] is not None else None
                                else:
                                    # 按中国时区(UTC+8)换算，
                                    # 与 split_db_by_day 的 CN 口径一致（原用本地时区，
                                    # 非 UTC+8 机器会整体偏移一天）。
                                    d_min = (datetime.fromtimestamp(row[0] / 1000, tz=CN)
                                             .strftime("%Y-%m-%d") if row[0] else None)
                                    d_max = (datetime.fromtimestamp(row[1] / 1000, tz=CN)
                                             .strftime("%Y-%m-%d") if row[1] else None)
                                if d_min and (not min_date or d_min < min_date):
                                    min_date = d_min
                                # 去掉 break——原实现在首个有数据的日期列后
                                # 即跳出，后续日期列不参与聚合，展示范围可能偏窄。
                                if d_max and (not max_date or d_max > max_date):
                                    max_date = d_max
                        except Exception:
                            # 收窄裸 except（该日期列缺失/不可解析时跳过）
                            continue
                except Exception:
                    # 单表统计失败不影响整体加载，收窄为 Exception
                    pass

            self.db_info = {
                "path": path,
                "tables": len(self.tables),
                "total_rows": total_rows,
                "min_date": min_date,
                "max_date": max_date,
                "table_info": table_info
            }

            # 显示信息
            info = f"数据库: {path}\n"
            info += f"表数量: {len(self.tables)} 张\n"
            info += f"总数据量: {total_rows:,} 行\n"
            if min_date and max_date:
                info += f"数据时间范围: {min_date} ~ {max_date}"

            self.info_text.config(state=tk.NORMAL)
            self.info_text.delete(1.0, tk.END)
            self.info_text.insert(1.0, info)
            self.info_text.config(state=tk.DISABLED)

            # 生成表选择列表
            self.generate_table_list()

            # 设置默认导出目录
            if not self.export_dir.get():
                self.export_dir.set(os.path.dirname(path))

            self.log(f"数据库加载成功: {len(self.tables)} 张表, {total_rows:,} 行数据")

        except Exception as e:
            messagebox.showerror("错误", f"加载数据库失败: {str(e)}")
            self.log(f"加载数据库失败: {str(e)}")
        finally:
            # 连接关闭移入 finally（异常路径不再泄漏）
            if conn is not None:
                try:
                    conn.close()
                except Exception:
                    pass

    def generate_table_list(self):
        """生成表选择列表"""
        # 清空现有内容
        for widget in self.table_inner.winfo_children():
            widget.destroy()

        self.table_vars = {}

        for i, (table, count) in enumerate(self.db_info["table_info"]):
            var = tk.BooleanVar(value=True)
            self.table_vars[table] = var

            frame = ttk.Frame(self.table_inner)
            frame.grid(row=i, column=0, sticky="w", padx=5, pady=2)

            ttk.Checkbutton(frame, variable=var, command=self.update_selected_count).pack(side=tk.LEFT)
            ttk.Label(frame, text=f"{table}  ({count:,} 行)", width=45, anchor="w").pack(side=tk.LEFT, padx=5)

        self.update_selected_count()

    def update_selected_count(self):
        """更新已选择表数量"""
        count = sum(1 for var in self.table_vars.values() if var.get())
        self.selected_count_label.config(text=f"（已选择 {count} 张表）")

    def select_all_tables(self):
        """全选"""
        for var in self.table_vars.values():
            var.set(True)
        self.update_selected_count()

    def deselect_all_tables(self):
        """全不选"""
        for var in self.table_vars.values():
            var.set(False)
        self.update_selected_count()

    def invert_tables(self):
        """反选"""
        for var in self.table_vars.values():
            var.set(not var.get())
        self.update_selected_count()

    def on_time_change(self):
        """时间范围改变"""
        pass

    def get_date_range(self):
        """获取选择的日期范围"""
        choice = self.time_range.get()

        if choice == "all":
            return None, None
        elif choice == "custom":
            start = self.custom_start.get().strip()
            end = self.custom_end.get().strip()
            if not start or not end:
                messagebox.showerror("错误", "请填写自定义日期范围！")
                return None, None
            # 预校验 8 位数字日期——非法输入此前会在
            # 导出线程的 strptime 处抛错，导致时间戳类表整批导出失败。
            for _label, _v in (("开始日期", start), ("结束日期", end)):
                try:
                    datetime.strptime(_v, "%Y%m%d")
                except ValueError:
                    messagebox.showerror("错误", f"{_label}格式应为 8 位数字 YYYYMMDD，当前为：{_v}")
                    return None, None
            if start > end:
                messagebox.showerror("错误", f"开始日期（{start}）不能晚于结束日期（{end}）")
                return None, None
            return start, end
        else:
            days = int(choice)
            # get_date_range 此前用 datetime.now()（本地时区），
            # 与同文件 L277-280 已统一的 CN(UTC+8) 口径混用，非 UTC+8 机器上时间窗口整体偏移一天。
            end_date = datetime.now(CN)
            start_date = end_date - timedelta(days=days)
            return start_date.strftime("%Y%m%d"), end_date.strftime("%Y%m%d")

    def start_export(self):
        """开始导出"""
        # 检查
        if not self.db_path.get() or not os.path.exists(self.db_path.get()):
            messagebox.showerror("错误", "请先选择并加载数据库！")
            return

        selected_tables = [t for t, var in self.table_vars.items() if var.get()]
        if not selected_tables:
            messagebox.showerror("错误", "请至少选择一张表！")
            return

        if not self.export_dir.get():
            messagebox.showerror("错误", "请选择导出目录！")
            return

        start_date, end_date = self.get_date_range()
        if self.time_range.get() == "custom" and (not start_date or not end_date):
            return

        # 禁用导出按钮
        self.export_btn.config(state=tk.DISABLED)
        self.progress["value"] = 0

        # 在新线程中执行导出
        # tk 变量（StringVar.get() 走 Tcl 解释器）不应跨线程
        # 访问——与 G1 已确立的"控件/弹窗经 _safe_after 回流"策略保持一致。
        # 改在**主线程**一次性取值后作为参数传入工作线程。
        thread = threading.Thread(
            target=self.do_export,
            args=(selected_tables, start_date, end_date,
                  self.export_format.get(), self.export_dir.get(), self.db_path.get()))
        thread.daemon = True
        thread.start()

    def do_export(self, selected_tables, start_date, end_date,
                  export_format, export_dir, db_path):
        """执行导出（export_format/export_dir/db_path 由主线程读取后传入，线程安全）"""
        conn = None  # G2修复：预置 None，保证 finally 中可安全关闭
        try:
            # 导出目录名仅秒级时间戳，同秒两次导出会因同名目录
            # 互相覆盖；追加微秒后缀保证唯一（与 auto_export_and_analyze 报告文件名同策略）。
            _ts = datetime.now(CN)
            timestamp = _ts.strftime("%Y%m%d_%H%M%S") + ("_%06d" % _ts.microsecond)

            # 创建导出子目录
            if export_format == "csv":
                out_dir = os.path.join(export_dir, f"health_export_{timestamp}")
                os.makedirs(out_dir, exist_ok=True)
            else:
                out_dir = export_dir

            self.log(f"开始导出...")
            self.log(f"格式: {export_format.upper()}")
            if start_date and end_date:
                self.log(f"时间范围: {start_date} ~ {end_date}")
            else:
                self.log(f"时间范围: 全部数据")
            self.log(f"选择表数: {len(selected_tables)}")

            # 连接数据库
            # db_path 已由主线程读取并作为参数传入，不再在此跨线程读 tk 变量
            conn = sqlite3.connect(db_path)
            conn.row_factory = sqlite3.Row

            total_tables = len(selected_tables)
            exported_tables = 0
            total_rows = 0

            # 用于 JSON/SQLite/Excel 的数据收集
            all_data = {}

            for i, table in enumerate(selected_tables):
                try:
                    self.log(f"正在处理 [{i+1}/{total_tables}]: {table}")

                    # 构建查询
                    query = f'SELECT * FROM {q_ident(table)}'
                    params = []

                    # 时间过滤
                    if start_date and end_date:
                        # 尝试不同的日期字段
                        date_cols = self.get_date_columns(conn, table)
                        if date_cols:
                            date_col = date_cols[0]
                            if date_col == "date":
                                query += f' WHERE date >= ? AND date <= ?'
                                params = [start_date, end_date]
                            else:
                                # 时间戳字段
                                # 补 CN(UTC+8) 后再换算。
                                # 原 naive .timestamp() 按宿主机本地时区解释，与本文件
                                # get_date_range 的 CN 口径（L425）不一致——用户选的
                                # 日期范围在非 UTC+8 机器上会整体偏移一天。
                                start_ts = int(datetime.strptime(start_date, "%Y%m%d").replace(tzinfo=CN).timestamp() * 1000)
                                end_ts = int((datetime.strptime(end_date, "%Y%m%d") + timedelta(days=1)).replace(tzinfo=CN).timestamp() * 1000)
                                query += f' WHERE {date_col} >= ? AND {date_col} < ?'
                                params = [start_ts, end_ts]

                    cursor = conn.execute(query, params)
                    columns = [desc[0] for desc in cursor.description]
                    rows = cursor.fetchall()

                    # 时间过滤后 0 行时区分两种情况（C-P2-11）：
                    #   ① 该表时间列全 0/无效（如 DBBreathRateStat）——过滤必然 0 行，
                    #      回退整表导出并告警；② 时间列有效、数据只是不在所选窗口内——
                    #      尊重过滤条件返回 0 行，不回退（旧行为会静默全量导出）。
                    #   params 非空 = 时间过滤确实生效过（无时间列时 params 为空，无此问题）。
                    if not rows and params:
                        cnt = conn.execute(f'SELECT COUNT(*) FROM {q_ident(table)}').fetchone()[0]
                        # _nz：时间列上有有效值（>0）的行数；探测失败按"无法判定"处理，
                        #   置 _nz=cnt（>0）→ 不满足回退条件，不回退（保守）。
                        try:
                            _nz = conn.execute(
                                f'SELECT COUNT(*) FROM {q_ident(table)} WHERE {q_ident(date_col)} > 0'
                            ).fetchone()[0]
                        except sqlite3.Error:
                            _nz = cnt
                        if cnt > 0 and _nz == 0:
                            self.log(f"  ⚠️  {table}: 时间范围内 0 行且时间列全 0/无效（整表 {cnt} 行），"
                                     f"回退整表导出")
                            query = f'SELECT * FROM {q_ident(table)}'
                            cursor = conn.execute(query)
                            columns = [desc[0] for desc in cursor.description]
                            rows = cursor.fetchall()

                    if not rows:
                        self.log(f"  ⚠️  {table}: 没有符合条件的数据，跳过")
                        continue

                    # 转换为字典列表
                    dict_rows = []
                    for row in rows:
                        row_dict = {}
                        for col in columns:
                            val = row[col]
                            if isinstance(val, bytes):
                                val = val.hex()
                            row_dict[col] = val
                        dict_rows.append(row_dict)

                    # CSV 分支逐表立即写文件，无需继续把全部行
                    # 留在 all_data 中——此前无论何种格式都全量驻留，峰值内存随库规模
                    # 线性增长；只有 JSON / SQLite / Excel 需要跨表聚合。
                    if export_format != "csv":
                        all_data[table] = {
                            "count": len(dict_rows),
                            "columns": columns,
                            "rows": dict_rows
                        }

                    # CSV 格式：每表一个文件
                    if export_format == "csv":
                        # C-P2-12：表名白名单化，防异常名目录穿越/错位（同 export 脚本）
                        csv_path = os.path.join(out_dir, f"{safe_fs_name(table)}.csv")
                        with atomic_write(csv_path, encoding="utf-8-sig", newline="") as f:
                            writer = csv.DictWriter(f, fieldnames=columns)
                            writer.writeheader()
                            writer.writerows(dict_rows)

                    total_rows += len(dict_rows)
                    exported_tables += 1
                    self.log(f"  ✅ {table}: {len(dict_rows):,} 行")

                except Exception as e:
                    self.log(f"  ❌ {table}: 导出失败 - {str(e)}")

                # 更新进度（G1修复：进度条属主线程控件，工作线程经 after 调度）
                pct = (i + 1) / total_tables * 100
                self._safe_after(0, lambda p=pct: self.progress.config(value=p))

            # 其他格式：写入单个文件
            if export_format == "json":
                out_path = os.path.join(out_dir, f"health_data_{timestamp}.json")
                with atomic_write(out_path, encoding="utf-8") as f:
                    json.dump(all_data, f, ensure_ascii=False, indent=2)
                self.log(f"JSON 文件已保存: {out_path}")

            elif export_format == "sqlite":
                out_path = os.path.join(out_dir, f"health_data_{timestamp}.db")
                if os.path.exists(out_path):
                    os.remove(out_path)
                # 连接关闭移入 finally——
                out_conn = sqlite3.connect(out_path)
                try:
                    for table, data in all_data.items():
                        cols = data["columns"]
                        col_defs = ", ".join([f'{q_ident(c)} TEXT' for c in cols])
                        out_conn.execute(f'CREATE TABLE {q_ident(table)} ({col_defs})')
                        placeholders = ", ".join(["?" for _ in cols])
                        col_names = ", ".join([f'{q_ident(c)}' for c in cols])
                        for row in data["rows"]:
                            values = [str(row[c]) if row[c] is not None else None for c in cols]
                            out_conn.execute(f'INSERT INTO {q_ident(table)} ({col_names}) VALUES ({placeholders})', values)
                    out_conn.commit()
                finally:
                    try:
                        out_conn.close()
                    except Exception:
                        pass
                # 明确告知与主工具类型化库的口径差异
                self.log(f"SQLite 文件已保存: {out_path}（各列均为 TEXT 类型；需要数值类型化库请运行 json_to_sqlite.py）")

            elif export_format == "excel":
                try:
                    from openpyxl import Workbook

                    # 所有表均无符合条件的数据时 all_data 为空，
                    # openpyxl 的 Workbook 至少需要 1 个 sheet，0 sheet 保存会失败。
                    if not all_data:
                        # 盲审复审 L8：空数据直接 return 会跳过 do_export 尾部的
                        #   "导出完成"日志与完成弹窗（CSV/JSON/SQLite 三路径都有收尾），
                        #   用户只看到一条 ⚠️ 易误判失败——此处补齐等价收尾再返回。
                        self.log("⚠️  没有符合条件的数据，Excel 导出结束（0 张表 / 0 行）")
                        self._safe_after(0, lambda: self.progress.config(value=100))
                        self._safe_after(0, lambda: messagebox.showinfo(
                            "完成", "没有符合条件的数据，Excel 导出结束（0 张表 / 0 行）。\n"
                                    "请检查所选时间范围或表后重试。"))
                        return

                    out_path = os.path.join(out_dir, f"health_data_{timestamp}.xlsx")
                    wb = Workbook()
                    wb.remove(wb.active)

                    for table, data in all_data.items():
                        # Excel sheet 名最长31字符；N4修复：
                        # 截断后可能撞名（如 DBBloodOxygenSaturationDataStat 与
                        # DBBloodOxygenSaturationDataStatTable 截断后同名），去重处理。
                        # C-P3-7：先 safe_fs_name 去掉 Excel sheet 名非法字符（:\/?*[] 等）再截断
                        sheet_name = safe_fs_name(table)[:31]
                        base_name = sheet_name
                        suffix = 2
                        existing = {ws.title for ws in wb.worksheets}
                        while sheet_name in existing:
                            tail = f"~{suffix}"
                            sheet_name = base_name[:31 - len(tail)] + tail
                            suffix += 1
                        ws = wb.create_sheet(title=sheet_name)

                        # 写入表头
                        ws.append(data["columns"])

                        # 写入数据
                        for row in data["rows"]:
                            ws.append([row[c] for c in data["columns"]])

                    wb.save(out_path)
                    self.log(f"Excel 文件已保存: {out_path}")

                except ImportError:
                    self.log("❌ 未安装 openpyxl，请执行: pip install openpyxl")
                    self._safe_after(0, lambda: messagebox.showerror("错误", "未安装 openpyxl 库，请执行: pip install openpyxl"))
                    return

            # 完成（G1修复：进度条与弹窗均调度回主线程）
            self._safe_after(0, lambda: self.progress.config(value=100))
            self.log("=" * 50)
            self.log(f"🎉 导出完成！")
            self.log(f"  导出表数: {exported_tables}")
            self.log(f"  总行数: {total_rows:,}")
            # GUI 全列导出补身份列隐私提醒，
            # 与命令行版（export_health_data.py 完成时告警）口径一致。
            self.log("⚠️ 隐私提醒：产物为全列导出，含 ssoid / open_id / device_unique_id 等身份列；"
                     "对外分享或投喂第三方 AI 前请先删列（做法见手册 9.2）")
            if export_format == "csv":
                self.log(f"  输出目录: {out_dir}")
            else:
                self.log(f"  输出文件: {out_path}")

            # 完成（G1修复：弹窗调度回主线程）
            self._safe_after(0, lambda: messagebox.showinfo(
                "完成", f"导出完成！\n\n导出表数: {exported_tables}\n总行数: {total_rows:,}\n输出目录: {out_dir}"
                        f"\n\n⚠️ 全列导出含身份列（ssoid 等），对外分享前请先删列（见手册 9.2）"))

        except Exception as e:
            self.log(f"❌ 导出失败: {str(e)}")
            self._safe_after(0, lambda err=str(e): messagebox.showerror("错误", f"导出失败: {err}"))
        finally:
            # 连接关闭移入 finally（此前异常路径会泄漏 sqlite 连接）；
            # 按钮恢复同样调度回主线程（finally 在工作线程中执行）。
            if conn is not None:
                try:
                    conn.close()
                except Exception:
                    pass
            self._safe_after(0, lambda: self.export_btn.config(state=tk.NORMAL))

    def get_date_columns(self, conn, table):
        """获取表中的日期列"""
        cursor = conn.execute(f'PRAGMA table_info({q_ident(table)})')
        columns = [row[1] for row in cursor.fetchall()]

        date_cols = []
        # 剔除 modified_*/update_*/created_* 等「元数据
        # 时间戳」——它们不是事件日，误用会把档案/配置表错误地按时间过滤。
        # 候选列改为与 split_db_by_day 共用同一份
        # utils.EVENT_TIME_COLUMNS（并保留 `date` 置顶以走 YYYYMMDD 精确匹配），
        for col in ["date"] + list(EVENT_TIME_COLUMNS):
            if col in columns:
                date_cols.append(col)

        return date_cols


def main():
    # 先确认图形环境可用，再创建窗口；失败时给出可执行的替代方案。
    if not _TK_AVAILABLE:
        print("❌ 无法启动图形界面：本机 Python 未提供 tkinter。")
        print(f"   原因：{_TK_ERROR}")
        print("   Windows：重装 Python 时勾选 “tcl/tk and IDLE”；")
        print("   Linux  ：安装 python3-tk（如 sudo apt install python3-tk）。")
        print("   或者改用命令行导出（功能等价）：python auto_export_and_analyze.py")
        return 1
    try:
        root = tk.Tk()
    except tk.TclError as e:
        print("❌ 无法启动图形界面：当前环境没有可用的显示（DISPLAY）。")
        print(f"   原因：{e}")
        print("   服务器 / CI / 容器等无显示环境请改用命令行：")
        print("       python auto_export_and_analyze.py")
        return 1

    # 设置主题
    try:
        style = ttk.Style()
        style.theme_use("clam")
    except tk.TclError:
        # 收窄为具体异常（指定主题不存在时抛 TclError），
        # 避免裸 except 掩盖其它真实错误。
        pass

    root.app = HealthExportApp(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())
